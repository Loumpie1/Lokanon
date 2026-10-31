"""Pilot-Paket zusammenstellen.

    python build\\pilot_paket.py [--version 0.2.0] [--installer <Pfad zur Setup.exe>]
                                 [--ziel %USERPROFILE%\\Lokanon\\Pilot] [--work %USERPROFILE%\\Lokanon\\build]

Legt unter <ziel> die Ordner 1-6 plus LIESMICH.txt an und packt alles außer dem Installer nach
<work>\\out\\Lokanon-Pilot-Paket-<version>.zip. Markdown-Unterlagen werden mit python-docx in Word-Dateien
übersetzt (Überschriften, Absätze, Listen, Tabellen, fett).
"""
from __future__ import annotations

import argparse
import hashlib
import os
import re
import shutil
import zipfile
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
WORK_DEFAULT = Path(os.environ.get("USERPROFILE", str(Path.home()))) / "Lokanon" / "build"
ZIEL_DEFAULT = Path(os.environ.get("USERPROFILE", str(Path.home()))) / "Lokanon" / "Pilot"


# ---------------------------------------------------------------- Markdown -> Word

def _runs(par, text: str) -> None:
    """**fett** und `code` als Runs; Links [text](url) -> text."""
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    for teil in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not teil:
            continue
        if teil.startswith("**") and teil.endswith("**"):
            par.add_run(teil[2:-2]).bold = True
        elif teil.startswith("`") and teil.endswith("`"):
            r = par.add_run(teil[1:-1])
            r.font.name = "Consolas"
        else:
            par.add_run(teil)


def md_nach_docx(md: str, ziel: Path, titel_fallback: str) -> None:
    import docx
    from docx.shared import Pt

    doc = docx.Document()
    doc.core_properties.author = "Lokanon"
    stil = doc.styles["Normal"]
    stil.font.name = "Calibri"
    stil.font.size = Pt(11)
    zeilen = md.splitlines()
    i = 0
    absatz: list[str] = []

    def absatz_schreiben():
        if absatz:
            _runs(doc.add_paragraph(), " ".join(absatz))
            absatz.clear()

    while i < len(zeilen):
        z = zeilen[i].rstrip()
        if not z.strip():
            absatz_schreiben()
            i += 1
            continue
        if z.startswith("#"):
            absatz_schreiben()
            ebene = len(z) - len(z.lstrip("#"))
            doc.add_heading(z.lstrip("#").strip(), level=min(ebene, 3) if ebene > 1 else 0)
            i += 1
            continue
        if z.strip() in ("---", "***"):
            absatz_schreiben()
            i += 1
            continue
        if z.lstrip().startswith(("- ", "* ")):
            absatz_schreiben()
            _runs(doc.add_paragraph(style="List Bullet"), z.lstrip()[2:].strip())
            i += 1
            continue
        if re.match(r"^\d+\.\s", z.lstrip()):
            absatz_schreiben()
            _runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\.\s", "", z.lstrip()))
            i += 1
            continue
        if z.startswith("|"):
            absatz_schreiben()
            rows = []
            while i < len(zeilen) and zeilen[i].startswith("|"):
                cells = [c.strip() for c in zeilen[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                breite = max(len(r) for r in rows)
                tab = doc.add_table(rows=len(rows), cols=breite)
                tab.style = "Table Grid"
                for ri, r in enumerate(rows):
                    for ci in range(breite):
                        zelle = tab.cell(ri, ci)
                        zelle.text = ""
                        _runs(zelle.paragraphs[0], r[ci] if ci < len(r) else "")
                        if ri == 0:
                            for run in zelle.paragraphs[0].runs:
                                run.bold = True
                doc.add_paragraph()
            continue
        absatz.append(z.strip())
        i += 1
    absatz_schreiben()
    if not doc.paragraphs:
        doc.add_heading(titel_fallback, 0)
    ziel.parent.mkdir(parents=True, exist_ok=True)
    doc.save(ziel)


# ---------------------------------------------------------------- Paket

LIESMICH = """LOKANON – PILOTPAKET
Pilotversion {version}, Stand {datum}

Lokanon ersetzt in Mandatsunterlagen alles, was eine Person erkennbar macht,
durch Platzhalter – lokal auf Ihrem PC, ohne Internet. Dateien oder einen
Ordner ablegen, prüfen, fertig: Zu jeder Datei entsteht eine Kopie im gleichen
Format mit dem Zusatz „_anonymisiert“ im Namen. Die Kopien können Sie in
ChatGPT oder Claude verwenden.


INHALT DIESES ORDNERS

  1 Installer\\            Lokanon-Setup-{version}.exe (ca. 1,4 GB) und SHA256.txt
  2 Anleitung\\            Lokanon-Anleitung.html – im Browser öffnen, läuft offline
                          (dieselbe Anleitung finden Sie im Programm unter „?“)
  3 Pilot-Unterlagen\\     Pilot-Hinweis (was Lokanon kann, was nicht, was wir zurückbrauchen)
                          und Rückmeldebogen – als Word-Datei und als Textfassung (.md)
  4 Rechtliche Vorlagen\\  Muster für Kanzleien: Übersicht, Checkliste Anbieterprüfung,
                          Verarbeitungsverzeichnis, Ergänzung der Datenschutzinformation
                          (Word-Dateien; vor Verwendung anwaltlich prüfen und anpassen)
  5 KI-Anweisung\\         Projekt-Anweisung zum Einfügen in ChatGPT/Claude (Textdatei)
                          und der Claude-Skill „lokanon-pseudonyme“ als ZIP
  6 Video\\                Lokanon-Teaser.mp4 – 30 Sekunden Produktvideo (ohne Ton)


INSTALLATION (ETWA 5 MINUTEN)

  1. „1 Installer\\Lokanon-Setup-{version}.exe“ doppelklicken.
  2. Windows SmartScreen meldet „Der Computer wurde durch Windows geschützt“,
     weil die Pilotversion noch nicht signiert ist. Klicken Sie auf
     „Weitere Informationen“ und dann auf „Trotzdem ausführen“.
  3. Dem Setup folgen. Es installiert nur für Ihr Windows-Benutzerkonto, braucht
     keine Administratorrechte und belegt etwa 2,8 GB. Internet ist nicht nötig,
     alle Sprachmodelle sind enthalten. Eine ältere Version wird ersetzt,
     Ihre Einstellungen bleiben erhalten.
  4. Lokanon über das Startmenü starten (Desktop-Verknüpfung im Setup wählbar).
     Der erste Start dauert ein bis zwei Minuten, weil die Sprachmodelle geladen
     werden; das Fenster zeigt solange „Lokanon wird gestartet …“.
  5. Beim ersten Start fragt ein Assistent nach Kanzleiname, Mitarbeitern und
     Kontaktdaten. Diese Angaben werden nie ersetzt.

  Die anonymisierten Kopien liegen bei „Ordner wählen“ neben den Originalen,
  sonst unter  C:\\Users\\<Benutzer>\\Dokumente\\Lokanon\\Anonymisiert
  Einstellungen und Protokoll:  C:\\Users\\<Benutzer>\\Lokanon
    einstellungen.yaml Ihre Einstellungen
    lokanon.log        Protokoll, falls Lokanon einmal nicht startet
    schluessel\\        verschlüsselte Schlüssel (nur mit eingeschalteter Rückübersetzung)
    Vorgaenge\\         Prüfbericht und Protokoll (nur mit eingeschaltetem Vorgangsordner)

  Deinstallation: Windows-Einstellungen → Apps → Lokanon → Deinstallieren.
  Der Datenordner C:\\Users\\<Benutzer>\\Lokanon bleibt dabei erhalten.

  Prüfsumme: Die Datei SHA256.txt enthält die SHA-256-Prüfsumme des Installers.
  In PowerShell prüfen mit:  Get-FileHash ".\\Lokanon-Setup-{version}.exe"


SO GEHT ES WEITER

  1. Anleitung lesen (2 Anleitung) – zehn Minuten.
  2. Pilot-Hinweis lesen (3 Pilot-Unterlagen).
  3. KI-Anbieter und Tarif klären (4 Rechtliche Vorlagen, Checkliste).
     Consumer-Konten wie ChatGPT Plus oder Claude Pro bieten keinen
     Auftragsverarbeitungsvertrag; geschäftliche Tarife (ChatGPT Business,
     Claude Team) schon.
  4. Projekt-Anweisung (5 KI-Anweisung) in Ihr ChatGPT-Projekt oder
     Claude-Projekt einfügen.
  5. Einen unkritischen Fall komplett durchspielen und den Rückmeldebogen
     ausfüllen.


GRENZEN DER PILOTVERSION

  - Keine Erkennung ist vollständig. Indirekte Hinweise (Rolle + Ort +
    Einrichtung) erkennt Lokanon nicht. Die Prüfung durch Sie ist Pflicht und
    fester Schritt im Programm.
  - Mit eingeschalteter Rückübersetzung sind die Daten pseudonymisiert, nicht
    anonym. Für die Kanzlei bleibt der Personenbezug bestehen.
  - Der Installer ist nicht signiert (SmartScreen-Hinweis, siehe oben).
  - Nur das Kanzlei-Profil ist enthalten, kein Profil für Arztpraxen.
  - Word- und Excel-Kopien bereinigen Text und Metadaten, aber noch keine
    Kommentare, Änderungsverfolgung oder Textfelder. Im Zweifel in den
    Einstellungen zusätzlich die Textfassung (.md) erzeugen lassen.
  - Scans werden per Texterkennung gelesen; schlechte Scans, Handschrift und
    Stempel können Fehler verursachen. PDF-Sichtkopien gibt Lokanon nur nach
    bestandener Gegenprüfung aus.
"""


def _leeren(p: Path) -> None:
    if p.exists():
        for kind in p.iterdir():
            if kind.is_dir():
                shutil.rmtree(kind)
            else:
                kind.unlink()
    p.mkdir(parents=True, exist_ok=True)


def paket(version: str, installer: Path | None, ziel: Path, work: Path) -> Path:
    ziel.mkdir(parents=True, exist_ok=True)
    # 1 Installer
    d1 = ziel / "1 Installer"
    d1.mkdir(exist_ok=True)
    if installer:
        for alt in d1.glob("Lokanon-Setup-*.exe"):
            alt.unlink()
        shutil.copy2(installer, d1 / installer.name)
        h = hashlib.sha256(installer.read_bytes()).hexdigest()
        (d1 / "SHA256.txt").write_text(f"{h}  {installer.name}\n", encoding="utf-8")
        (work / "out").mkdir(parents=True, exist_ok=True)
        (work / "out" / f"{installer.name}.sha256").write_text(f"{h}  {installer.name}\n", encoding="utf-8")
    # 2 Anleitung
    d2 = ziel / "2 Anleitung"
    _leeren(d2)
    html = (ROOT / "static" / "anleitung" / "index.html").read_text(encoding="utf-8")
    (d2 / "Lokanon-Anleitung.html").write_text(html.replace('href="/ui/fonts.css"', 'href="fonts.css"'), encoding="utf-8")
    shutil.copy2(ROOT / "ui" / "fonts.css", d2 / "fonts.css")
    shutil.copytree(ROOT / "ui" / "fonts", d2 / "fonts")
    # 3 Pilot-Unterlagen
    d3 = ziel / "3 Pilot-Unterlagen"
    _leeren(d3)
    for quelle, name in (("Pilot-Hinweis.md", "Pilot-Hinweis"), ("Rueckmeldebogen.md", "Rückmeldebogen")):
        md = (ROOT / "docs" / "pilot" / quelle).read_text(encoding="utf-8")
        (d3 / f"{name}.md").write_text(md, encoding="utf-8")
        md_nach_docx(md, d3 / f"{name}.docx", name)
    # 4 Rechtliche Vorlagen
    d4 = ziel / "4 Rechtliche Vorlagen"
    _leeren(d4)
    for quelle, name in (("LIESMICH.md", "0 LIESMICH – Übersicht"),
                         ("checkliste-anbieterpruefung.md", "1 Checkliste Anbieterprüfung"),
                         ("verarbeitungsverzeichnis-eintrag.md", "2 Verarbeitungsverzeichnis – Mustereintrag"),
                         ("art13-ergaenzung.md", "3 Art. 13 DSGVO – Ergänzung Datenschutzinformation")):
        md = (ROOT / "docs" / "vorlagen" / quelle).read_text(encoding="utf-8")
        md_nach_docx(md, d4 / f"{name}.docx", name)
    # 5 KI-Anweisung
    d5 = ziel / "5 KI-Anweisung"
    _leeren(d5)
    shutil.copy2(ROOT / "ki-anweisung" / "projekt-anweisung.md", d5 / "Projekt-Anweisung fuer ChatGPT oder Claude.txt")
    skill = ROOT / "ki-anweisung" / "lokanon-pseudonyme"
    with zipfile.ZipFile(d5 / "lokanon-pseudonyme-skill.zip", "w", zipfile.ZIP_DEFLATED) as z:
        for f in sorted(skill.rglob("*")):
            if f.is_file():
                z.write(f, f"lokanon-pseudonyme/{f.relative_to(skill).as_posix()}")
    # 6 Video
    d6 = ziel / "6 Video"
    _leeren(d6)
    teaser = ROOT / "video" / "out" / "lokanon-teaser.mp4"
    if teaser.exists():
        shutil.copy2(teaser, d6 / "Lokanon-Teaser.mp4")
    # LIESMICH
    (ziel / "LIESMICH.txt").write_text(
        LIESMICH.format(version=version, datum=date.today().strftime("%d.%m.%Y")), encoding="utf-8")
    # ZIP ohne Installer
    out = work / "out" / f"Lokanon-Pilot-Paket-{version}.zip"
    out.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for f in sorted(ziel.rglob("*")):
            if f.is_file() and "1 Installer" not in f.parts:
                z.write(f, f"Lokanon-Pilot-Paket-{version}/{f.relative_to(ziel).as_posix()}")
    return out


def main() -> None:
    ap = argparse.ArgumentParser(description="Lokanon Pilot-Paket zusammenstellen")
    ap.add_argument("--version", default=(ROOT / "build" / "VERSION").read_text(encoding="utf-8").strip())
    ap.add_argument("--installer", type=Path, help="Setup.exe, die nach '1 Installer' kopiert wird")
    ap.add_argument("--ziel", type=Path, default=ZIEL_DEFAULT)
    ap.add_argument("--work", type=Path, default=WORK_DEFAULT)
    a = ap.parse_args()
    installer = a.installer
    if installer is None:
        kandidat = a.work / "out" / f"Lokanon-Setup-{a.version}.exe"
        installer = kandidat if kandidat.exists() else None
    out = paket(a.version, installer, a.ziel, a.work)
    print(f"Paket: {a.ziel}")
    print(f"ZIP:   {out} ({out.stat().st_size / 1024 / 1024:.1f} MB)")
    if installer:
        print(f"Installer: {installer}")
    else:
        print("Kein Installer gefunden – Ordner '1 Installer' unverändert.")


if __name__ == "__main__":
    main()
