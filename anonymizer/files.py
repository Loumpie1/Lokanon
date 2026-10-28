"""Dateien lesen und anonymisierte Kopien schreiben. Originale werden nie verändert.

PDF-Sichtkopie (write_pdf / verify_pdf)
---------------------------------------
Die Originale werden auf Zeichenebene gesucht (get_text("rawdict"): Groß-/Klein-
schreibung egal, Whitespace flexibel, auch über Zeilenumbrüche und Trennstriche hinweg),
als Redaction-Annotationen entfernt (apply_redactions: Text raus, überlappende Bild-
teile ausgeweißt) und der Platzhalter danach mit insert_textbox in das breiteste
Teilrechteck geschrieben. Passt er nicht, wird die Schrift bis 5 pt verkleinert, dann
die Kurzform ([PERSON_2] -> [P2]) genommen, notfalls über den Rand geschrieben; nach dem
Schreiben wird per get_text geprüft, dass jeder Platzhalter sichtbar ist (sonst ValueError).
Anschließend: alle Annotationen und Formularfelder (inkl. AcroForm) löschen, Lesezeichen
leeren, doc.scrub() (Metadaten, XMP, Anhänge, JavaScript, Links, Thumbnails, Seiten
bereinigen), Speichern mit garbage=4/deflate/clean, nie inkrementell.
Zum Schluss läuft verify_pdf über Seitentext, alle Objekte und Streams (dekomprimiert)
und Metadaten; ein Fund ergibt ValueError. Treffersuche, Fallback-Prüfung und verify_pdf
benutzen dieselben Wortgrenzen wie _replace_all ("Engel" trifft nicht "Engelmann"), damit
die Sichtkopie zum Text-/Markdown-Ergebnis passt.

MuPDF-Workaround (PyMuPDF 2.0 / MuPDF 1.28.5): Entfernt apply_redactions einen Lauf, der
mitten in einem String-Element eines TJ-Arrays beginnt, fehlt im Ausgleich die Breite des
ersten entfernten Glyphs und der Folgetext der Zeile rutscht nach links unter das Füll-
rechteck (bei Tj-Operatoren stimmt alles). write_pdf schreibt deshalb vorher alle TJ-Arrays
der Seiteninhalte und Form-XObjects so um, dass jedes Element genau einen Glyph-Code
enthält (1 Byte bei einfachen Schriften, 2 Byte bei Type0/Identity-H; andere CMaps werden
nicht angefasst) und prüft je Seite, dass Zeichen und Positionen unverändert sind.
Außerdem wird /helv in die eigenen Schriftressourcen der Seite eingetragen, weil PyMuPDF
sonst eine gleichnamige XObject-Schrift findet und der Platzhalter in einer Ersatzschrift
gerendert würde.

Form-XObjects: PyMuPDF 2.0 (MuPDF) filtert beim apply_redactions auch in Form-XObjects
rekursiv – getestet mit page.show_pdf_page (tests/test_pdf.py, testdaten/pdf/
make_xobject_pdf.py): der Text im XObject wird entfernt, auch der Stream selbst enthält
ihn nicht mehr. Als Sicherheitsnetz gibt es trotzdem einen Fallback: bleibt auf einer
Seite nach dem Apply ein Original im Seitentext stehen, wird die Seite gerastert
(200 dpi) und als Bild neu eingefügt; die alte Seite verschwindet mit garbage=4.
Unsichtbarer Text (render mode 3, Deckkraft 0, OCR-Schicht "GlyphLessFont") wird im
selben Redaction-Durchlauf entfernt (Kriterien wie Document.scrub); scrub läuft danach mit
hidden_text=False, weil sein eigener Durchlauf sonst Platzhalter-Glyphen neben ehemaligen
Hidden-Text-Flächen abschneiden kann. Nicht abgedeckt: Text, der nur als
Grafikpfade oder in Bildern (Scans) vorliegt – dafür ist OCR nötig.
"""
from __future__ import annotations

import csv
import io
from datetime import datetime


def _doc_key(data: bytes) -> str:
    import hashlib

    return hashlib.sha1(data).hexdigest()


def read_text(name: str, data: bytes) -> str:
    return read_document(name, data)[0]


def read_document(name: str, data: bytes) -> tuple[str, dict]:
    """Text einer Datei plus Infos: {"seiten": n, "ocr_seiten": n} bei PDF.
    Gescannte PDF-Seiten (ohne Textebene) werden automatisch per OCR gelesen."""
    ext = name.lower().rsplit(".", 1)[-1]
    info: dict = {}
    if ext in ("txt", "md"):
        return data.decode("utf-8", errors="replace"), info
    if ext == "csv":
        text = data.decode("utf-8-sig", errors="replace")
        rows = csv.reader(io.StringIO(text), delimiter=";" if text.count(";") > text.count(",") else ",")
        return "\n".join(" | ".join(r) for r in rows), info
    if ext == "docx":
        import docx

        d = docx.Document(io.BytesIO(data))
        return "\n".join(p.text for p in _docx_paragraphs(d)), info
    if ext == "pdf":
        import pymupdf as fitz

        from .ocr import ocr_text

        with fitz.open(stream=data, filetype="pdf") as pdf:
            if pdf.needs_pass:
                raise ValueError("PDF ist passwortgeschützt.")
            texte, n_ocr = ocr_text(pdf, key=_doc_key(data))
            info = {"seiten": len(pdf), "ocr_seiten": n_ocr}
        text = "\n".join(texte)
        if not text.strip():
            raise ValueError("PDF enthält keinen erkennbaren Text (leere Seiten oder unleserlicher Scan).")
        return text, info
    if ext == "xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        lines = []
        for ws in wb.worksheets:
            lines.append(f"## {ws.title}")
            for row in ws.iter_rows(values_only=True):
                vals = [str(v) for v in row if v is not None]
                if vals:
                    lines.append(" | ".join(vals))
        return "\n".join(lines), info
    raise ValueError(f"Dateityp .{ext} wird nicht unterstützt.")


def _docx_paragraphs(d):
    yield from d.paragraphs
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for s in d.sections:
        for part in (s.header, s.footer, s.first_page_header, s.first_page_footer):
            yield from part.paragraphs


def _replace_all(s: str, pairs) -> str:
    import re

    for orig, ph in pairs:
        s = re.sub(rf"(?<!\w){re.escape(orig)}(?!\w)", ph, s, flags=re.IGNORECASE)
    return s


def write_docx(data: bytes, pairs) -> bytes:
    """Ersetzt Begriffe in einer DOCX-Kopie. Formatierung bleibt erhalten, außer ein
    Begriff ist über mehrere Formatierungs-Runs verteilt (dann wird der Absatz vereinfacht)."""
    import docx

    d = docx.Document(io.BytesIO(data))
    for p in _docx_paragraphs(d):
        if not p.text:
            continue
        for run in p.runs:
            run.text = _replace_all(run.text, pairs)
        if _replace_all(p.text, pairs) != p.text:  # Rest über Run-Grenzen
            p.text = _replace_all(p.text, pairs)
    _strip_docx_meta(d)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _strip_docx_meta(d):
    cp = d.core_properties
    for attr in ("author", "last_modified_by", "title", "subject", "keywords", "comments",
                 "category", "content_status", "identifier", "language", "version"):
        setattr(cp, attr, "")
    cp.created = cp.modified = datetime(2000, 1, 1)
    cp.revision = 1


def write_xlsx(data: bytes, pairs) -> bytes:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and not c.value.startswith("="):
                    c.value = _replace_all(c.value, pairs)
        ws.title = _replace_all(ws.title, pairs).replace("[", "").replace("]", "")[:31]
    p = wb.properties
    p.creator = p.lastModifiedBy = p.title = p.subject = p.keywords = p.description = None
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF: geschwärzte Sichtkopie
# ---------------------------------------------------------------------------

_PDF_MIN_FONTSIZE = 5.0


def _pdf_norm(s: str) -> str:
    """Kleinbuchstaben, Whitespace-Folgen (auch Zeilenumbrüche) zu einem Leerzeichen."""
    import re

    return re.sub(r"\s+", " ", s).strip().lower()


def _pdf_leak_regex(orig: str):
    """Original als ganzes Wort im normalisierten Text (gleiche Wortgrenzen wie die
    Treffersuche und _replace_all: 'Engel' trifft nicht 'Engelmann')."""
    import re

    return re.compile(r"(?<!\w)" + re.escape(_pdf_norm(orig)) + r"(?!\w)")


def _pdf_short_placeholder(ph: str) -> str:
    """[PERSON_12] -> [P12]; alles andere wird nur abgeschnitten."""
    import re

    m = re.fullmatch(r"\[([A-Za-z])[A-Za-z_]*?_?(\d+)\]", ph)
    if m:
        return f"[{m.group(1).upper()}{m.group(2)}]"
    inner = ph.strip("[]")
    return f"[{inner[:3].upper()}]" if inner else ph


def _pdf_page_chars(page):
    """Alle Zeichen der Seite als Liste (zeichen, bbox, fontsize, zeilen-id) plus den
    daraus gebildeten Seitentext (Zeilen durch Zeilenumbruch getrennt). Index im Text ==
    Index in der Zeichenliste, Zeilenumbrüche haben bbox None."""
    import pymupdf

    flags = pymupdf.TEXT_PRESERVE_WHITESPACE | pymupdf.TEXT_MEDIABOX_CLIP
    chars: list = []
    for bi, block in enumerate(page.get_text("rawdict", flags=flags)["blocks"]):
        for li, line in enumerate(block.get("lines", [])):
            for span in line.get("spans", []):
                size = float(span.get("size") or 10)
                for ch in span.get("chars", []):
                    chars.append((ch["c"], pymupdf.Rect(ch["bbox"]), size, (bi, li)))
            chars.append(("\n", None, 0.0, (bi, li)))
    return chars, "".join(c[0] for c in chars)


def _pdf_needle_regex(orig: str):
    """Groß-/Kleinschreibung egal, Whitespace im Original matcht beliebigen Whitespace
    (auch Zeilenumbruch), zwischen zwei Zeichen darf ein Trennstrich-Umbruch stehen, und
    wie in _replace_all gelten Wortgrenzen (kein Treffer mitten im Wort)."""
    import re

    parts = []
    for tok in re.split(r"(\s+)", orig.strip()):
        if not tok:
            continue
        if tok.isspace():
            parts.append(r"\s+")
        else:
            parts.append(r"(?:-\n)?".join(re.escape(c) for c in tok))
    return re.compile(r"(?<!\w)" + "".join(parts) + r"(?!\w)", re.IGNORECASE)


def _pdf_find_hits(page, pairs):
    """Liefert je Treffer (platzhalter, [rect je Zeile], fontsize). Längere Originale
    zuerst; ein Treffer innerhalb einer schon belegten Fläche wird übersprungen."""
    import pymupdf

    chars, text = _pdf_page_chars(page)
    hits = []
    claimed: list = []
    for orig, ph in pairs:
        if not orig or not orig.strip():
            continue
        rx = _pdf_needle_regex(orig)
        for m in rx.finditer(text):
            per_line: dict = {}
            size = 0.0
            for c, bbox, fs, lid in chars[m.start():m.end()]:
                if bbox is None or bbox.is_empty:
                    continue
                per_line[lid] = per_line[lid] | bbox if lid in per_line else pymupdf.Rect(bbox)
                size = max(size, fs)
            rects = [r for r in per_line.values() if not r.is_empty]
            if not rects:
                continue
            area = sum(r.get_area() for r in rects)
            covered = sum((r & c).get_area() for r in rects for c in claimed)
            if covered > 0.5 * area:
                continue  # liegt in einem längeren, schon markierten Treffer
            claimed.extend(rects)
            hits.append((ph, rects, size or 10.0))
    return hits


def _pdf_draw_placeholder(page, rect, ph: str, size: float) -> tuple:
    """Schreibt den Platzhalter in das Rechteck. Erst Schrift verkleinern (bis 5 pt),
    dann Kurzform, zuletzt ungeachtet der Breite einfügen - nie stillschweigend weglassen.
    Gibt (geschriebener Text, Rechteck, Schriftgröße) zurück."""
    import pymupdf

    size = max(min(size, rect.height / 1.15), _PDF_MIN_FONTSIZE)
    for text in (ph, _pdf_short_placeholder(ph)):
        fs = size
        while fs >= _PDF_MIN_FONTSIZE:
            width = pymupdf.get_text_length(text, "helv", fs)
            if width <= rect.width - 1:
                box = pymupdf.Rect(rect)
                box.y1 = max(box.y1, box.y0 + fs * 1.25)
                box.x1 = max(box.x1, box.x0 + width + 2)
                rc = page.insert_textbox(box, text, fontname="helv", fontsize=fs, align=1)
                if rc >= 0:
                    return text, box, fs
            fs -= 0.5
    text = _pdf_short_placeholder(ph)
    page.insert_text((rect.x0, rect.y1 - 1), text, fontname="helv", fontsize=_PDF_MIN_FONTSIZE)
    return text, pymupdf.Rect(rect), _PDF_MIN_FONTSIZE


def _pdf_count(haystack: str, needle: str) -> int:
    return haystack.replace("\n", "").replace(" ", "").count(needle)


_PDF_WS = b"\x00\t\n\x0c\r "
_PDF_DELIM = b"()<>[]{}/%"


def _pdf_literal_end(data: bytes, i: int) -> int:
    """Ende (exklusiv) des Literal-Strings, der bei data[i] == '(' beginnt."""
    n = len(data)
    depth, j = 1, i + 1
    while j < n and depth:
        ch = data[j:j + 1]
        if ch == b"\\":
            j += 2
            continue
        if ch == b"(":
            depth += 1
        elif ch == b")":
            depth -= 1
        j += 1
    return j


def _pdf_scan(data: bytes):
    """Grober Tokenizer für Content-Streams: liefert (art, start, ende). Arrays und
    Dictionaries werden als ein Token geliefert, Inline-Bilder (BI..ID..EI) als 'op'."""
    import re

    num = re.compile(rb"^[+-]?(?:\d+\.?\d*|\.\d+)$")
    ei = re.compile(rb"(?:^|\s)EI(?=\s|$)")
    toks = []
    i, n = 0, len(data)
    stop = _PDF_WS + _PDF_DELIM
    while i < n:
        c = data[i:i + 1]
        if c in _PDF_WS:
            j = i + 1
            while j < n and data[j:j + 1] in _PDF_WS:
                j += 1
            toks.append(("ws", i, j))
        elif c == b"%":
            j = i + 1
            while j < n and data[j:j + 1] not in b"\r\n":
                j += 1
            toks.append(("ws", i, j))
        elif c == b"(":
            toks.append(("str", i, _pdf_literal_end(data, i)))
        elif c == b"<":
            if data[i + 1:i + 2] == b"<":
                depth, j = 1, i + 2
                while j < n and depth:
                    if data[j:j + 2] == b"<<":
                        depth += 1
                        j += 2
                    elif data[j:j + 2] == b">>":
                        depth -= 1
                        j += 2
                    elif data[j:j + 1] == b"(":
                        j = _pdf_literal_end(data, j)
                    else:
                        j += 1
                toks.append(("dict", i, j))
            else:
                j = data.find(b">", i + 1)
                j = n if j < 0 else j + 1
                toks.append(("hex", i, j))
        elif c == b"[":
            depth, j = 1, i + 1
            while j < n and depth:
                ch = data[j:j + 1]
                if ch == b"[":
                    depth += 1
                    j += 1
                elif ch == b"]":
                    depth -= 1
                    j += 1
                elif ch == b"(":
                    j = _pdf_literal_end(data, j)
                elif ch == b"<" and data[j + 1:j + 2] != b"<":
                    k = data.find(b">", j + 1)
                    j = n if k < 0 else k + 1
                else:
                    j += 1
            toks.append(("array", i, j))
        elif c == b"/":
            j = i + 1
            while j < n and data[j:j + 1] not in stop:
                j += 1
            toks.append(("name", i, j))
        elif c in b"]})>{":
            toks.append(("other", i, i + 1))
        else:
            j = i
            while j < n and data[j:j + 1] not in stop:
                j += 1
            if j == i:
                j = i + 1
            tok = data[i:j]
            if tok == b"ID":  # Inline-Bild: Binärdaten bis EI überspringen
                m = ei.search(data, j)
                j = n if not m else m.end()
                toks.append(("op", i, j))
            else:
                toks.append(("num" if num.match(tok) else "op", i, j))
        i = toks[-1][2]
    return toks


def _pdf_literal_bytes(raw: bytes) -> bytes:
    """Dekodiert einen PDF-Literal-String '(...)' zu Bytes."""
    s = raw[1:-1]
    esc = {0x6E: 10, 0x72: 13, 0x74: 9, 0x62: 8, 0x66: 12}
    out = bytearray()
    i, n = 0, len(s)
    while i < n:
        c = s[i]
        if c == 0x5C:  # Backslash
            i += 1
            if i >= n:
                break
            c = s[i]
            if c in esc:
                out.append(esc[c])
                i += 1
            elif c == 0x0D:
                i += 1 + (s[i + 1:i + 2] == b"\n")
            elif c == 0x0A:
                i += 1
            elif 0x30 <= c <= 0x37:
                j = i
                while j < n and j < i + 3 and 0x30 <= s[j] <= 0x37:
                    j += 1
                out.append(int(s[i:j], 8) & 0xFF)
                i = j
            else:
                out.append(c)
                i += 1
        elif c == 0x0D:
            out.append(10)
            i += 1 + (s[i + 1:i + 2] == b"\n")
        else:
            out.append(c)
            i += 1
    return bytes(out)


def _pdf_split_tj_array(raw: bytes, code_len: int) -> bytes:
    """Zerlegt alle Strings eines TJ-Arrays in Elemente mit je einem Glyph-Code."""
    inner = raw[1:-1]
    out = [b"["]
    for kind, s, e in _pdf_scan(inner):
        tok = inner[s:e]
        if kind == "str":
            b = _pdf_literal_bytes(tok)
        elif kind == "hex":
            h = bytes(ch for ch in tok[1:-1] if ch not in _PDF_WS)
            if len(h) % 2:
                h += b"0"
            try:
                b = bytes.fromhex(h.decode("ascii"))
            except ValueError:
                out.append(tok)
                continue
        else:
            out.append(tok)
            continue
        if len(b) <= code_len:
            out.append(tok)
            continue
        for k in range(0, len(b), code_len):
            out.append(b"<" + b[k:k + code_len].hex().encode() + b">")
    out.append(b"]")
    return b"".join(out)


def _pdf_split_tj_stream(data: bytes, fonts: dict) -> bytes:
    """Schreibt jedes '[...] TJ' so um, dass jedes Array-Element genau einen Glyph-Code
    enthält (nur für Schriften mit bekannter Code-Länge). Sonst bleibt der Stream gleich."""
    toks = _pdf_scan(data)
    out = []
    operands: list = []
    font_len = None
    stack: list = []
    changed = False
    for kind, s, e in toks:
        piece = data[s:e]
        if kind == "op":
            if piece == b"Tf" and operands and operands[0][0] == "name":
                font_len = fonts.get(operands[0][1][1:])
            elif piece == b"q":
                stack.append(font_len)
            elif piece == b"Q":
                font_len = stack.pop() if stack else font_len
            elif piece == b"TJ" and font_len and operands and operands[-1][0] == "array":
                idx, raw = operands[-1][2], operands[-1][1]
                new = _pdf_split_tj_array(raw, font_len)
                if new != raw:
                    out[idx] = new
                    changed = True
            operands = []
        elif kind != "ws":
            operands.append((kind, piece, len(out)))
        out.append(piece)
    return b"".join(out) if changed else data


def _pdf_code_len(ftype: str, enc: str):
    if ftype in ("Type1", "TrueType", "Type3", "MMType1"):
        return 1
    if ftype == "Type0" and enc in ("Identity-H", "Identity-V"):
        return 2
    return None


def _pdf_snapshot(page):
    return [(c, round(b.x0, 1), round(b.y0, 1)) for c, b, _fs, _l in _pdf_page_chars(page)[0] if b]


def _pdf_split_tj(doc) -> None:
    """Workaround für MuPDF 1.28 (PyMuPDF 2.0): Beim Entfernen eines Textstücks, das
    mitten in einem TJ-String-Element beginnt, verrechnet sich apply_redactions um die
    Breite des ersten entfernten Glyphs, der Folgetext rutscht nach links. Mit Ein-Glyph-
    Elementen tritt das nicht auf. Jede Seite wird nach dem Umbau geprüft (Zeichen und
    Positionen unverändert), sonst werden die Streams zurückgesetzt."""
    done: set = set()
    for pno in range(doc.page_count):
        page = doc[pno]
        fonts: dict = {}
        try:
            for xref, _ext, ftype, _base, name, enc, ref in page.get_fonts(full=True):
                fonts.setdefault(0 if ref == 0 else ref, {})[name.encode("latin-1", "replace")] = \
                    _pdf_code_len(ftype, enc)
            streams = [(x, 0) for x in page.get_contents()]
            streams += [(x, x) for x, _n, _inv, _bb in page.get_xobjects()]
        except Exception:
            continue
        before = _pdf_snapshot(page)
        changed = []
        for sx, owner in streams:
            if sx in done or not fonts.get(owner):
                continue
            done.add(sx)
            try:
                old = doc.xref_stream(sx)
                if not old:
                    continue
                new = _pdf_split_tj_stream(old, fonts[owner])
                if new != old:
                    doc.update_stream(sx, new)
                    changed.append((sx, old))
            except Exception:
                continue
        if changed:
            try:
                page = doc.reload_page(page)
            except Exception:
                page = doc[pno]
            if _pdf_snapshot(page) != before:
                for sx, old in changed:
                    doc.update_stream(sx, old)


def _pdf_hidden_rects(page) -> list:
    """Rechtecke aller unsichtbaren Text-Spans (gleiche Kriterien wie Document.scrub:
    GlyphLessFont/OCR-Schicht, Deckkraft 0, weder gefüllt noch konturiert)."""
    import pymupdf

    filled = getattr(pymupdf.mupdf, "FZ_STEXT_FILLED", None)
    stroked = getattr(pymupdf.mupdf, "FZ_STEXT_STROKED", None)
    flags = pymupdf.TEXT_PRESERVE_SPANS | getattr(pymupdf, "TEXT_COLLECT_STYLES", 0)
    rects = []
    for block in page.get_text("dict", flags=flags).get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                font = span.get("font") or ""
                hidden = font.split("+")[-1] == "GlyphLessFont" or span.get("alpha") == 0
                cf = span.get("char_flags")
                if isinstance(cf, int) and isinstance(filled, int) and isinstance(stroked, int):
                    hidden = hidden or not (cf & filled or cf & stroked)
                if hidden:
                    r = pymupdf.Rect(span["bbox"])
                    if not r.is_empty and not r.is_infinite:
                        rects.append(r)
    return rects


def _pdf_ensure_font(page) -> None:
    """Stellt sicher, dass /helv in den eigenen Ressourcen der Seite steht. PyMuPDF
    findet sonst eine gleichnamige Schrift aus einem XObject und trägt nichts ein; der
    Platzhalter würde dann mit einer Ersatzschrift gerendert."""
    doc = page.parent
    xref = page.insert_font(fontname="helv")
    if doc.xref_get_key(page.xref, "Resources/Font/helv")[0] != "null":
        return
    cur, prefix = page.xref, []
    for key in ("Resources", "Font"):
        t, v = doc.xref_get_key(cur, "/".join(prefix + [key]))
        if t == "xref":
            cur, prefix = int(v.split()[0]), []
        else:
            prefix.append(key)
    doc.xref_set_key(cur, "/".join(prefix + ["helv"]), f"{xref} 0 R")


def write_pdf(data: bytes, pairs: list[tuple[str, str]]) -> bytes:
    """Geschwärzte PDF-Sichtkopie: Originale werden als Redaction entfernt, der Platz-
    halter an ihre Stelle geschrieben, alles Beiwerk (Metadaten, Anhänge, Formulare,
    Lesezeichen, Links, JavaScript, versteckter Text) entfernt. Am Ende prüft
    verify_pdf das Ergebnis; findet es noch ein Original, gibt es ValueError."""
    import pymupdf

    pairs = sorted(((o, p) for o, p in pairs if o and len(o.strip()) >= 3), key=lambda t: -len(t[0]))
    originals = [o for o, _ in pairs]
    leak_rx = [_pdf_leak_regex(o) for o in originals]
    doc = pymupdf.open(stream=data, filetype="pdf")
    try:
        if doc.needs_pass:
            raise ValueError("PDF ist passwortgeschützt und kann nicht geschwärzt werden.")
        _pdf_split_tj(doc)
        rasterize: list[int] = []
        drawn: dict[int, list] = {}
        for page in doc:
            for w in list(page.widgets()):
                page.delete_widget(w)
            hits = _pdf_find_hits(page, pairs)
            for _ph, rects, _size in hits:
                for r in rects:
                    page.add_redact_annot(r, fill=(1, 1, 1))
            hidden = _pdf_hidden_rects(page)  # unsichtbarer Text (OCR-Schicht, Tr 3, Alpha 0)
            for r in hidden:
                page.add_redact_annot(r, cross_out=False)
            if hits or hidden:
                page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_PIXELS,
                                      graphics=pymupdf.PDF_REDACT_LINE_ART_REMOVE_IF_COVERED,
                                      text=pymupdf.PDF_REDACT_TEXT_REMOVE)
            while page.first_annot:
                page.delete_annot(page.first_annot)
            expected: dict[str, int] = {}
            if hits:
                _pdf_ensure_font(page)
            for ph, rects, size in hits:
                widest = max(rects, key=lambda r: r.width)
                written, box, fs = _pdf_draw_placeholder(page, widest, ph, size)
                drawn.setdefault(page.number, []).append((written, box, fs))
                expected[written] = expected.get(written, 0) + 1
            page_text = page.get_text()
            missing = [t for t, n in expected.items() if _pdf_count(page_text, t) < n]
            if missing:
                raise ValueError(f"Platzhalter nicht sichtbar auf Seite {page.number + 1}: {missing}")
            norm = _pdf_norm(page_text)
            if any(rx.search(norm) for rx in leak_rx):
                rasterize.append(page.number)
        # Scan-Seiten (ohne Textebene): OCR + Pixel-Schwärzung + Gegenprüfung (anonymizer/ocr.py)
        from .ocr import schwaerze_scanseiten

        schwaerze_scanseiten(doc, pairs, key=_doc_key(data))
        # Fallback: Seiten, auf denen Text die Redaction überlebt hat (z. B. exotische
        # XObjects), werden gerastert und als Bild neu eingefügt.
        for pno in sorted(rasterize, reverse=True):
            old = doc[pno]
            pix = old.get_pixmap(dpi=200, annots=False)
            new = doc.new_page(pno=pno + 1, width=old.rect.width, height=old.rect.height)
            new.insert_image(new.rect, pixmap=pix)
            for text, box, fs in drawn.get(pno, []):  # Platzhalter bleiben extrahierbar
                new.insert_textbox(box, text, fontname="helv", fontsize=fs, align=1)
            doc.delete_page(pno)
        doc.set_toc([])
        try:
            doc.xref_set_key(doc.pdf_catalog(), "AcroForm", "null")
        except Exception:
            pass
        # hidden_text=False: unsichtbarer Text wurde oben schon mit denselben Kriterien
        # entfernt; ein zweiter Durchlauf könnte Platzhalter-Glyphen neben alten
        # Hidden-Text-Flächen mit abschneiden.
        doc.scrub(attached_files=True, clean_pages=True, embedded_files=True, hidden_text=False,
                  javascript=True, metadata=True, redactions=True, remove_links=True,
                  reset_fields=True, reset_responses=True, thumbnails=True, xml_metadata=True)
        doc.set_metadata({})
        out = doc.tobytes(garbage=4, deflate=True, clean=True)
    finally:
        doc.close()
    leaked = verify_pdf(out, originals)
    if leaked:
        raise ValueError(f"PDF enthält nach der Schwärzung noch Originaltext: {leaked}")
    return out


def verify_pdf(data: bytes, originals: list[str]) -> list[str]:
    """Sucht jedes Original (Groß-/Kleinschreibung egal, Whitespace normalisiert, als
    ganzes Wort wie in _replace_all) im Seitentext, in allen Objekten/Streams
    (dekomprimiert) und in den Metadaten. Rückgabe: noch gefundene Originale; leer = sauber."""
    import pymupdf

    needles = [(o, _pdf_leak_regex(o)) for o in originals if o and len(o.strip()) >= 3]
    if not needles:
        return []
    doc = pymupdf.open(stream=data, filetype="pdf")
    try:
        parts = [page.get_text() for page in doc]
        parts.extend(str(v) for v in (doc.metadata or {}).values() if v)
        try:
            parts.append(doc.get_xml_metadata() or "")
        except Exception:
            pass
        for x in range(1, doc.xref_length()):
            for getter in (doc.xref_object, doc.xref_stream):
                try:
                    raw = getter(x)
                except Exception:
                    continue
                if not raw:
                    continue
                if isinstance(raw, bytes):
                    parts.append(raw.decode("latin-1"))
                    parts.append(raw.decode("utf-16", errors="ignore"))
                else:
                    parts.append(str(raw))
    finally:
        doc.close()
    hay = _pdf_norm("\n".join(parts))
    seen: set = set()
    found = []
    for orig, rx in needles:
        if rx.pattern not in seen and rx.search(hay):
            seen.add(rx.pattern)
            found.append(orig)
    return found
