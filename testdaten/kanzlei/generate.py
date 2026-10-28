"""Erzeugt fünf fiktive Kanzlei-Fälle (Erb- und Verkehrsrecht) als Testdokumente
plus Referenzmarkierungen (JSONL) für evaluate.py.

Markup im Text: «LABEL|Wert»  -> Wert ist sensibel und muss maskiert werden.
Alle Personen, Kontakte, Konten, Kennzeichen und Nummern sind erfunden.
Echt sind nur Straßennamen, Städte und öffentliche Gerichte.
"""
import json
import re
import sys
from pathlib import Path

import docx
import openpyxl
from docx.shared import Pt
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))
from anonymizer.files import read_text  # noqa: E402

OUT = Path(__file__).resolve().parent
OUT.mkdir(parents=True, exist_ok=True)
MARK = re.compile(r"«([A-Z_]+)\|([^»]+)»")


def iban(blz, konto):
    """Prüfziffernkorrekte deutsche IBAN (fiktives Konto)."""
    bban = f"{blz}{konto:010d}"
    n = int(bban + "131400")  # D=13, E=14, 00
    check = 98 - n % 97
    raw = f"DE{check:02d}{bban}"
    return " ".join(raw[i:i + 4] for i in range(0, len(raw), 4))


def render(template):
    """Markup entfernen, Spans im Klartext merken."""
    out, spans, pos = [], [], 0
    for m in MARK.finditer(template):
        out.append(template[pos:m.start()])
        start = sum(len(x) for x in out)
        out.append(m.group(2))
        spans.append([start, start + len(m.group(2)), m.group(1)])
        pos = m.end()
    out.append(template[pos:])
    return "".join(out), spans


def spans_by_search(text, template):
    """Für PDF/XLSX: markierte Werte im extrahierten Text suchen."""
    spans = []
    for lab, val in {(m.group(1), m.group(2)) for m in MARK.finditer(template)}:
        for m in re.finditer(rf"(?<![\w-]){re.escape(val)}(?![\w-])", text):
            spans.append([m.start(), m.end(), lab])
    # Teiltreffer innerhalb längerer Markierungen entfernen ("Engel" in "Markus Engel")
    spans = [s for s in spans if not any(o != s and o[0] <= s[0] and o[1] >= s[1] for o in spans)]
    return sorted(spans)


gold = []


def add_gold(fname, text, spans):
    gold.append({"datei": fname, "text": text, "spans": spans})


# ------------------------------------------------------------------ Fall 1
IBAN1 = iban("36250000", 7712045)
fall1 = f"""Kanzlei Lehmkuhl & Partner – Rechtsanwälte
Mandatsnotiz Erbrecht – Testamentsanfechtung

Mandantin: «PERSON|Sabine Winter» geb. «PERSON|Brinkhoff», geb. am «DATUM|09.06.1971»
Anschrift: «ADRESSE|Eppinghofer Straße 47», «PLZ_ORT|45468 Mülheim an der Ruhr»
Telefon: «TELEFON|0208 4471236», E-Mail: «EMAIL_ADDRESS|sabine.winter71@example.com»

Erblasser: «PERSON|Gerhard Brinkhoff», geb. «DATUM|14.02.1941», verstorben am «DATUM|03.08.2026»
zuletzt wohnhaft «ADRESSE|Mendener Straße 112», «PLZ_ORT|45470 Mülheim an der Ruhr»

Nachlassgericht: Amtsgericht Mülheim an der Ruhr, Georgstraße 13, 45468 Mülheim an der Ruhr
Aktenzeichen: «AKTENZEICHEN|14 VI 512/26»

Sachverhalt:
Frau «PERSON|Winter» ist die Tochter des Erblassers. Mit notariellem Testament vom 12.05.2024
(UR-Nr. «AKTENZEICHEN|318/2024» des Notars «PERSON|Dr. Stefan Lindemann», Essen) setzte Herr «PERSON|Brinkhoff»
seine frühere Pflegekraft «PERSON|Natalia Kowalczyk» als Alleinerbin ein. Der Bruder der Mandantin,
«PERSON|Thomas Brinkhoff», wurde ebenso wie die Mandantin enterbt.

Die Mandantin bezweifelt die Testierfähigkeit. Laut Hausarzt «PERSON|Dr. Michael Essen» lag bereits
im Frühjahr 2024 eine fortgeschrittene Erkrankung vor. Herr «PERSON|Essen» ist bereit, als Zeuge auszusagen.
Der Erblasser war bis 2019 «INDIREKT|langjähriger Schützenkönig in Saarn» und im Stadtteil bekannt.

Nachlass (vorläufig):
– Einfamilienhaus, Grundbuch von «GRUNDBUCH|Speldorf Blatt 4711», Flur 3, Flurstück 218
– Girokonto Sparkasse Mülheim, IBAN «IBAN_CODE|{IBAN1}»
– Pkw Mercedes E 220, amtl. Kennzeichen «KFZ|MH-GB 141»

Frau «PERSON|Kowalczyk» hat beim Amtsgericht Essen, Zweigertstraße 52, 45130 Essen, keinen Antrag gestellt;
der Erbscheinsantrag liegt in Mülheim. Rückruf Frau «PERSON|Winter» erbeten bis Freitag.
"""

# ------------------------------------------------------------------ Fall 2
IBAN2 = iban("36050105", 2201889)
IBAN2b = iban("36050105", 9100423)
fall2_brief = f"""# Pflichtteilsverlangen

**Lehmkuhl & Partner Rechtsanwälte**

Herrn
«PERSON|Jan Haverkamp»
«ADRESSE|Rellinghauser Straße 88»
«PLZ_ORT|45128 Essen»

Essen, 21.10.2026

**Nachlass «PERSON|Margarete Haverkamp», verstorben am «DATUM|17.06.2026»**
**Unser Zeichen: «AKTENZEICHEN|LP-2026/0417»**

Sehr geehrter Herr «PERSON|Haverkamp»,

wir vertreten die rechtlichen Interessen Ihrer Schwester, Frau «PERSON|Laura Haverkamp»,
«ADRESSE|Holsterhauser Straße 31», «PLZ_ORT|45147 Essen». Eine Vollmacht liegt bei.

Ihre Mutter hat Sie mit Testament vom 02.03.2021 zum Alleinerben eingesetzt. Unsere Mandantin
ist als Abkömmling pflichtteilsberechtigt (§ 2303 BGB). Wir fordern Sie auf, bis zum
15.11.2026 ein notarielles Nachlassverzeichnis vorzulegen (§ 2314 BGB).

Nach den uns vorliegenden Informationen umfasst der Nachlass insbesondere:

- Eigentumswohnung «ADRESSE|Rüttenscheider Straße 201», «PLZ_ORT|45131 Essen»
- Girokonto, IBAN «IBAN_CODE|{IBAN2}»
- Wertpapierdepot Nr. «KONTONUMMER|448 120 773»
- Pkw VW Tiguan, Kennzeichen «KFZ|E-MH 2048»

Die Steuer-ID der Erblasserin lautet «STEUER_ID|47 036 892 816», Steuernummer
«STEUERNUMMER|112/5034/1234» (Finanzamt Essen-Süd).

Zahlungen leisten Sie bitte ausschließlich auf unser Anderkonto
IBAN «IBAN_CODE|{IBAN2b}».

Für Rückfragen erreichen Sie mich unter 0201 8945120 oder per E-Mail an
kanzlei@lehmkuhl-partner.example.

Mit freundlichen Grüßen
Rechtsanwältin Dr. Annette Lehmkuhl
"""

fall2_xlsx_rows = [
    ("Position", "Beschreibung", "Wert (EUR)", "Bemerkung"),
    ("1", "Eigentumswohnung «ADRESSE|Rüttenscheider Straße 201», «PLZ_ORT|45131 Essen»", 285000, "Gutachten «PERSON|Ralf Overbeck» liegt vor"),
    ("2", f"Girokonto IBAN «IBAN_CODE|{IBAN2}»", 18432.17, "Stand Todestag"),
    ("3", "Depot Nr. «KONTONUMMER|448 120 773»", 64210.00, "Kurswert 17.06.2026"),
    ("4", "Pkw VW Tiguan «KFZ|E-MH 2048»", 14500, "Schätzwert DAT"),
    ("5", "Schmuck und Hausrat", 3500, "Angabe «PERSON|Jan Haverkamp»"),
    ("6", "Darlehen an Nachbarin «PERSON|Elke Sommer»", 5000, "nicht verbrieft"),
    ("7", "Beerdigungskosten", -7800, "Bestatter «ORGANIZATION|Pietät Kemper GmbH»"),
]

# ------------------------------------------------------------------ Fall 3
fall3 = """Kanzlei Lehmkuhl & Partner – Verkehrsrecht
Mandantenaufnahme Verkehrsunfall

Mandant: «PERSON|Kevin Schulte», geb. «DATUM|22.09.1993»
Anschrift: «ADRESSE|Duisburger Straße 290», «PLZ_ORT|45478 Mülheim an der Ruhr»
Mobil: «TELEFON|0176 23894451»
E-Mail: «EMAIL_ADDRESS|k.schulte93@example.com»
Rechtsschutz: ARAG, Versicherungsschein-Nr. «VERSICHERUNGSNR|RS-4417-882031»

Unfalltag: 14.10.2026, ca. 07:45 Uhr
Unfallort: Kreuzung Aktienstraße / Kaiserstraße, Mülheim an der Ruhr

Eigenes Fahrzeug: VW Golf VIII, amtl. Kennzeichen «KFZ|MH-KS 1993»
FIN: «FIN|WVWZZZCDZMW123456»

Unfallgegner: «PERSON|Murat Yildiz», «ADRESSE|Wanheimer Straße 15», «PLZ_ORT|47053 Duisburg»
Fahrzeug: BMW 320d, amtl. Kennzeichen «KFZ|DU-MY 777»
Haftpflicht: HUK-COBURG, Schadennummer «VERSICHERUNGSNR|SN 26-0912-55413»

Zeugin: «PERSON|Petra Bach», «ADRESSE|Kaiserstraße 9», «PLZ_ORT|45468 Mülheim an der Ruhr», Tel. «TELEFON|0208 381177»
Polizeiliche Unfallaufnahme: PI Mülheim, Vorgangsnummer «AKTENZEICHEN|VU 2026/14385»

Schilderung des Mandanten:
Herr «PERSON|Schulte» fuhr bei Grün in die Kreuzung ein. Herr «PERSON|Yildiz» bog aus der
Gegenrichtung nach links ab und übersah ihn. Frau «PERSON|Bach» stand an der Fußgängerampel
und hat den Unfall gesehen. Herr «PERSON|Schulte» hat Prellungen an der Schulter, Behandlung
im Evangelischen Krankenhaus Mülheim. Arbeitgeber: «ORGANIZATION|Spedition Rehbein GmbH», Ausfall 5 Arbeitstage.

Reparatur: «ORGANIZATION|Autohaus Rehbein & Söhne», Kostenvoranschlag 6.840,00 EUR netto.
Bankverbindung für Regulierung: IBAN «IBAN_CODE|{IBAN3}»
"""
fall3 = fall3.replace("{IBAN3}", iban("36250000", 3311902))

# ------------------------------------------------------------------ Fall 4 (PDF)
fall4 = [
    ("h", "Einspruch gegen den Bußgeldbescheid"),
    ("p", "Lehmkuhl & Partner Rechtsanwälte · Essen"),
    ("p", "An die Bußgeldstelle der Stadt Essen<br/>Ihr Aktenzeichen: «AKTENZEICHEN|7.1-552318/26»"),
    ("p", "Betroffene: «PERSON|Dr. Claudia Engel», geb. «DATUM|05.11.1979»<br/>"
          "«ADRESSE|Moltkestraße 15», «PLZ_ORT|45138 Essen»"),
    ("p", "Führerschein-Nr.: «FUEHRERSCHEIN|J010B7XY351»<br/>Fahrzeug: Audi A4, Kennzeichen «KFZ|E-CE 511»"),
    ("p", "Sehr geehrte Damen und Herren,"),
    ("p", "namens und in Vollmacht unserer Mandantin legen wir gegen den Bußgeldbescheid vom 06.10.2026 "
          "Einspruch ein. Vorgeworfen wird eine Überschreitung der zulässigen Höchstgeschwindigkeit um "
          "31 km/h auf der A 40 in Fahrtrichtung Dortmund, Höhe Anschlussstelle Essen-Frohnhausen."),
    ("p", "Frau «PERSON|Engel» bestreitet, zum Tatzeitpunkt Fahrerin gewesen zu sein. Das Fahrzeug wird auch von "
          "ihrem Ehemann «PERSON|Markus Engel» genutzt. Auf dem Messfoto ist eine männliche Person erkennbar. "
          "Wir beantragen Akteneinsicht einschließlich Messprotokoll, Eichschein und Rohmessdaten."),
    ("p", "Unsere Mandantin ist als «INDIREKT|Chefärztin der Kinderklinik am Elisabeth-Krankenhaus» beruflich auf "
          "die Fahrerlaubnis angewiesen; ein Fahrverbot hätte erhebliche Folgen."),
    ("p", "Mit freundlichen Grüßen<br/>Dr. Annette Lehmkuhl, Rechtsanwältin"),
]

# ------------------------------------------------------------------ Fall 5 (E-Mail-Verlauf)
fall5 = """Von: «EMAIL_ADDRESS|wolfgang.rhein@example.com»
An: kanzlei@lehmkuhl-partner.example
Datum: 24.10.2026 18:12
Betreff: Erbe Tante Hildegard und der Unfall mit dem Anhänger

Guten Abend Frau Dr. Lehmkuhl,

wie am Telefon besprochen, hier die Infos. Meine Tante «PERSON|Hildegard Rhein» ist am «DATUM|02.09.2026»
im Seniorenstift Haus Ruhrblick verstorben. Sie hat mir und meiner Cousine «PERSON|Jutta Lange» je die
Hälfte vermacht. Jutta will jetzt aber das Haus in der «ADRESSE|Mellinghofer Straße 64» allein behalten
und hat mir über ihren Anwalt «PERSON|Herrn Frank Delius» einen Abfindungsvorschlag geschickt.

Dann noch die andere Sache: Beim Ausräumen des Hauses bin ich mit dem geliehenen Anhänger
(«KFZ|MH-W 4210») gegen das Tor des Nachbarn «PERSON|Heinz-Dieter Sommer» gefahren. Herr «PERSON|Sommer»
will 2.300 Euro. Er ist übrigens «INDIREKT|der Vorsitzende des Kleingartenvereins Styrum-Nord}», falls das hilft.

Ich bin unter «TELEFON|+49 172 5583917» erreichbar, meine Frau «PERSON|Ingrid Rhein» unter «TELEFON|0208 593421».
Wir wohnen «ADRESSE|Hingbergstraße 230», «PLZ_ORT|45472 Mülheim an der Ruhr».

Viele Grüße
«PERSON|Wolle» («PERSON|Wolfgang Rhein»)

-----
Von: kanzlei@lehmkuhl-partner.example
An: «EMAIL_ADDRESS|wolfgang.rhein@example.com»
Datum: 25.10.2026 09:03

Sehr geehrter Herr «PERSON|Rhein»,

danke für die Unterlagen. Bitte schicken Sie mir noch den Erbschein und das Schreiben von
Herrn «PERSON|Delius». Termin am Donnerstag im Rhein-Ruhr-Zentrum klappt leider nicht,
gern aber in der Kanzlei.

Freundliche Grüße
A. Lehmkuhl
"""
fall5 = fall5.replace("}»", "»")


# ================================================================== Dateien schreiben
def write_docx(path, template, title_lines=2):
    text, spans = render(template.strip("\n"))
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    d.core_properties.author = "Dr. Annette Lehmkuhl"          # absichtlich: Metadaten-Test
    d.core_properties.title = "Mandat " + MARK.search(template).group(2)
    for i, line in enumerate(text.split("\n")):
        p = d.add_paragraph(line)
        if i < title_lines:
            p.runs and setattr(p.runs[0], "bold", True)
    d.save(path)
    extracted = read_text(path.name, path.read_bytes())
    assert extracted.rstrip("\n") == text, f"Extraktion weicht ab: {path.name}"
    add_gold(path.name, extracted, spans)


def write_plain(path, template):
    text, spans = render(template)
    path.write_text(text, encoding="utf-8")
    assert read_text(path.name, path.read_bytes()) == text
    add_gold(path.name, text, spans)


def write_xlsx(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Nachlassverzeichnis"
    wb.properties.creator = "Dr. Annette Lehmkuhl"
    template_all = []
    for r, row in enumerate(rows, 1):
        for c, val in enumerate(row, 1):
            if isinstance(val, str):
                template_all.append(val)
                val, _ = render(val)
            cell = ws.cell(r, c, val)
            if r == 1:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F5FA8")
            if c == 3 and r > 1:
                cell.number_format = "#,##0.00"
    for col, w in zip("ABCD", (10, 62, 14, 40)):
        ws.column_dimensions[col].width = w
    ws.cell(len(rows) + 1, 2, "Summe").font = Font(bold=True)
    ws.cell(len(rows) + 1, 3, f"=SUM(C2:C{len(rows)})").number_format = "#,##0.00"
    wb.save(path)
    extracted = read_text(path.name, path.read_bytes())
    add_gold(path.name, extracted, spans_by_search(extracted, "\n".join(template_all)))


def write_pdf(path, blocks):
    styles = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=styles["Normal"], fontName="Helvetica", fontSize=10.5, leading=15, spaceAfter=8)
    head = ParagraphStyle("h", parent=styles["Heading2"], fontName="Helvetica-Bold")
    doc = SimpleDocTemplate(str(path), pagesize=A4, author="Dr. Annette Lehmkuhl",
                            title="Einspruch Engel", leftMargin=60, rightMargin=60)
    story, template_all = [], []
    for kind, tpl in blocks:
        template_all.append(tpl)
        txt, _ = render(tpl)
        txt = txt.replace("&", "&amp;")
        story += [Paragraph(txt, head if kind == "h" else body), Spacer(1, 4)]
    doc.build(story)
    extracted = read_text(path.name, path.read_bytes())
    add_gold(path.name, extracted, spans_by_search(extracted, "\n".join(template_all)))


write_docx(OUT / "Fall1_Erbrecht_Testamentsanfechtung_Brinkhoff.docx", fall1)
write_plain(OUT / "Fall2_Erbrecht_Pflichtteilsverlangen_Haverkamp.md", fall2_brief)
write_xlsx(OUT / "Fall2_Erbrecht_Nachlassverzeichnis_Haverkamp.xlsx", fall2_xlsx_rows)
write_docx(OUT / "Fall3_Verkehrsrecht_Unfallaufnahme_Schulte.docx", fall3)
write_pdf(OUT / "Fall4_Verkehrsrecht_Einspruch_Bussgeld_Engel.pdf", fall4)
write_plain(OUT / "Fall5_Mischfall_EMail-Verlauf_Rhein.txt", fall5)

with open(OUT / "referenz_markierungen.jsonl", "w", encoding="utf-8") as fh:
    for g in gold:
        fh.write(json.dumps(g, ensure_ascii=False) + "\n")

for g in gold:
    print(f"{g['datei']:<55} {len(g['spans']):>3} markierte Stellen")
